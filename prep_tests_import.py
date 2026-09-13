# -*- coding: utf-8 -*-
"""Prepare les donnees de test pour l'import xlsx :
1. extrait l'AOA (tableau de tableaux) + les fusions du fichier tab_service_TCSLHF1.xlsx
   -> _aoa_service.json (simule ce que SheetJS renverrait dans l'app)
2. genere _eleves_test.xlsx (2 classes, format reel arab e) puis son AOA -> _aoa_eleves.json
"""
import io, json
from openpyxl import Workbook, load_workbook

def aoa_de(chemin):
    wb = load_workbook(chemin)
    sortie = {}
    for nom in wb.sheetnames:
        ws = wb[nom]
        aoa = []
        for ligne in ws.iter_rows(values_only=True):
            aoa.append(['' if v is None else str(v) for v in ligne])
        fusions = [{'s': {'r': m.min_row - 1, 'c': m.min_col - 1}, 'e': {'r': m.max_row - 1, 'c': m.max_col - 1}}
                   for m in ws.merged_cells.ranges]
        sortie[nom] = {'aoa': aoa, 'merges': fusions}
    return sortie

# 1. tableau de service de reference
data = aoa_de('tab_service_TCSLHF1.xlsx')
io.open('_aoa_service.json', 'w', encoding='utf-8').write(json.dumps(data, ensure_ascii=False))
print('_aoa_service.json :', {k: (len(v['aoa']), len(v['merges'])) for k, v in data.items()})

# 2. tableaux des eleves (2 classes, format arabe reel)
wb = Workbook()
def feuille(ws, classe, eleves):
    ws.append(['القسم', classe])
    ws.append([])
    ws.append(['رمز التلميذ', 'الاسم الكامل', 'الاسم بالحروف اللاتينية'])
    for code, ar, fr in eleves:
        ws.append([code, ar, fr])

ws1 = wb.active
ws1.title = '3A'
feuille(ws1, 'TCSLHF1', [
    ('J130012345', 'بنعلي فاطمة الزهراء', 'Benali Fatima'),
    ('J130012346', 'العمري أحمد', 'El Amrani Ahmed'),
    ('J130012347', 'علمي يوسف', 'Alami Youssef'),
    ('J130012348', 'الشرايبي نادية', 'Chraibi Nadia'),
])
ws2 = wb.create_sheet('3B')
feuille(ws2, 'TCSF-4', [
    ('J130012349', 'برادة سعاد', 'Berrada Souad'),
    ('J130012350', 'الطازي كريم', 'Tazi Karim'),
])
ws2.append(['ملاحظة', 'اختبار'])
wb.save('_eleves_test.xlsx')

data2 = aoa_de('_eleves_test.xlsx')
io.open('_aoa_eleves.json', 'w', encoding='utf-8').write(json.dumps(data2, ensure_ascii=False))
print('_aoa_eleves.json :', {k: len(v['aoa']) for k, v in data2.items()})
for k, v in data2.items():
    print('---', k); [print('   ', r) for r in v['aoa']]
