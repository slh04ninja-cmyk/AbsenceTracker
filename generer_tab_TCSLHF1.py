# -*- coding: utf-8 -*-
"""Genere le tableau de service de la classe TCSLHF1 en xlsx (mise en forme arabe, feuille RTL).

Structure reproduite de la page 1 du PDF de l'etablissement :
- en-tete administratif (royaume / ministere / direction / etablissement / annee / titre)
- grille : lignes = jours (lundi -> samedi), colonnes = 8 creneaux horaires, une matiere par case
  (cellules fusionnees pour les seances de 2 h), la salle en 2e ligne en petit
- liste des professeurs affectes par matiere (3 paires matiere/professeur, comme l'original)
"""
import io
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
except ImportError as e:
    raise SystemExit('openpyxl manquant : ' + str(e))

SORTIE = 'tab_service_TCSLHF1.xlsx'

CRENEAUX = ['08:30-09:30', '09:30-10:30', '10:30-11:30', '11:30-12:30',
            '14:00-15:00', '15:00-16:00', '16:00-17:00', '17:00-18:00']
JOURS = ['الإثنين', 'الثلاثاء', 'الأربعاء', 'الخميس', 'الجمعة', 'السبت']

# matiere, salle, colonne de debut (1 = 08:30-09:30 ... 8 = 17:00-18:00), nombre de creneaux
GRILLE = {
    'الإثنين': [('الرياضيات', 'Salle Lycée 3', 1, 1), ('التربية البدنية', 'بدون قاعة', 2, 1),
                ('التاريخ والجغرافيا', 'Salle Lycée 2', 3, 2), ('اللغة الفرنسية', 'Salle Lycée 3', 5, 1),
                ('اللغة الإنجليزية', 'Salle Lycée 2', 6, 1)],
    'الثلاثاء': [('اللغة الفرنسية', 'Salle Lycée 3', 1, 1), ('اللغة الإنجليزية', 'Salle Lycée 2', 2, 1),
                 ('اللغة العربية', 'Salle Lycée 5', 3, 1), ('علوم الحياة والأرض', 'Salle Lycée 1', 4, 1),
                 ('الرياضيات', 'Salle Lycée 3', 5, 1), ('التاريخ والجغرافيا', 'Salle Lycée 2', 6, 1),
                 ('الفلسفة', 'Salle Lycée 5', 7, 1)],
    'الأربعاء': [('اللغة العربية', 'Salle Lycée 5', 1, 1), ('اللغة الإنجليزية', 'Salle Lycée 2', 2, 1),
                 ('التربية الإسلامية', 'Salle Lycée 5', 6, 1), ('الفلسفة', 'Salle Lycée 3', 7, 1)],
    'الخميس': [('التربية الإسلامية', 'Salle Lycée 5', 1, 1), ('التاريخ والجغرافيا', 'Salle Lycée 2', 2, 1),
               ('اللغة العربية', 'Salle Lycée 5', 5, 2), ('اللغة الفرنسية', 'Salle Lycée 3', 7, 1)],
    'الجمعة': [('اللغة الإنجليزية', 'Salle Lycée 2', 1, 1), ('التربية البدنية', 'بدون قاعة', 2, 1),
               ('اللغة الفرنسية', 'Salle Lycée 3', 3, 2)],
    'السبت': []
}

PROFS = [('الرياضيات', 'أيوب الكمرة'), ('اللغة العربية', 'المهدي الصلحي'), ('اللغة الفرنسية', 'سامية الحاضي'),
         ('علوم الحياة والأرض', 'كمال الوردي'), ('الفلسفة', 'زكية المندريلي'), ('التربية البدنية', 'يسرى البوسعيدي'),
         ('التربية الإسلامية', 'محمد خليفي'), ('اللغة الإنجليزية', 'هشام أجامي'), ('التاريخ والجغرافيا', 'ياسين القامة')]

FIN = Side(style='thin', color='000000')
BORD = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)
GRIS = PatternFill('solid', fgColor='F2F2F2')
GRIS_FONCE = PatternFill('solid', fgColor='E7E6E6')
CENTRE = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb = Workbook()
ws = wb.active
ws.title = 'TCSLHF1'
ws.sheet_view.rightToLeft = True        # feuille en lecture droite -> gauche

LARG = 9
def ligne_entete(idx, texte, gras=True, taille=11, fond=None):
    ws.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=LARG)
    c = ws.cell(row=idx, column=1, value=texte)
    c.font = Font(bold=gras, size=taille)
    c.alignment = CENTRE
    if fond is not None:
        for col in range(1, LARG + 1):
            ws.cell(row=idx, column=col).fill = fond
    ws.row_dimensions[idx].height = 20
    return idx + 1

r = 1
r = ligne_entete(r, 'المملكة المغربية', True, 12)
r = ligne_entete(r, 'وزارة التربية الوطنية والتعليم الأولي والرياضة', False, 10)
r = ligne_entete(r, 'المديرية الإقليمية – سطات', False, 10)
r = ligne_entete(r, 'الثانوية القصبية الإعدادية – بوعوان', True, 11)
r = ligne_entete(r, 'الموسم الدراسي 2026-2027', False, 10)
r = ligne_entete(r, 'جدول حصص القسم : TCSLHF1', True, 13, GRIS)
r = ligne_entete(r, 'EDT for FET timetables — 06/09/2026 17:41', False, 8)
r += 1

# ---- grille ----
LIGNE_EN_TETE = r
ws.cell(row=r, column=1, value='اليوم / الحصة').font = Font(bold=True, size=10)
ws.cell(row=r, column=1).alignment = CENTRE
ws.cell(row=r, column=1).fill = GRIS_FONCE
for i, cr in enumerate(CRENEAUX):
    c = ws.cell(row=r, column=2 + i, value=cr)
    c.font = Font(bold=True, size=9)
    c.alignment = CENTRE
    c.fill = GRIS_FONCE
for col in range(1, LARG + 1):
    ws.cell(row=r, column=col).border = BORD
ws.row_dimensions[r].height = 24
r += 1

for jour in JOURS:
    c = ws.cell(row=r, column=1, value=jour)
    c.font = Font(bold=True, size=11)
    c.alignment = CENTRE
    c.fill = GRIS
    c.border = BORD
    for col in range(2, LARG + 1):
        cc = ws.cell(row=r, column=col)
        cc.border = BORD
        cc.alignment = CENTRE
        cc.font = Font(size=10)
    for matiere, salle, debut, duree in GRILLE[jour]:
        col = 1 + debut
        cellule = ws.cell(row=r, column=col)
        try:
            cellule.value = CellRichText(matiere, TextBlock(InlineFont(sz=8, color='595959'), '\n' + salle))
        except Exception:
            cellule.value = matiere + '\n' + salle
        if duree > 1:
            ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + duree - 1)
    ws.row_dimensions[r].height = 38
    r += 1

r += 1

# ---- liste des professeurs (3 paires matiere/professeur par ligne, comme la page d'origine) ----
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=LARG)
c = ws.cell(row=r, column=1, value='لائحة الأساتذة المسندين حسب المواد')
c.font = Font(bold=True, size=11)
c.alignment = CENTRE
for col in range(1, LARG + 1):
    ws.cell(row=r, column=col).fill = GRIS
    ws.cell(row=r, column=col).border = BORD
ws.row_dimensions[r].height = 20
r += 1

PAIRES = [1, 4, 7]           # colonnes de depart des 3 paires (RTL : la 1re paire est la plus a droite)
for ligne_idx in range(3):
    for i, col_dep in enumerate(PAIRES):
        matiere, prof = PROFS[ligne_idx + 3 * i]
        cm = ws.cell(row=r, column=col_dep, value=matiere)
        cp = ws.cell(row=r, column=col_dep + 1, value=prof)
        for cc in (cm, cp):
            cc.border = BORD
            cc.alignment = CENTRE
            cc.font = Font(size=10, bold=(cc is cm))
    ws.row_dimensions[r].height = 20
    r += 1

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=LARG)
c = ws.cell(row=r, column=1, value='استخراج آلي من ملف PDF — يُرجى التحقق من المعطيات  /  Extraction automatique — à valider')
c.font = Font(size=8, italic=True, color='808080')
c.alignment = CENTRE

# ---- mise en page ----
ws.column_dimensions['A'].width = 16
for i in range(2, LARG + 1):
    ws.column_dimensions[get_column_letter(i)].width = 13
ws.freeze_panes = ws.cell(row=LIGNE_EN_TETE + 1, column=2)
ws.page_setup.orientation = 'landscape'
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(SORTIE)
print('ecrit :', SORTIE)

# ---- relecture de controle ----
from openpyxl import load_workbook
wb2 = load_workbook(SORTIE)
ws2 = wb2['TCSLHF1']
print('feuille            :', ws2.title, '| dimensions :', ws2.dimensions, '| RTL :', ws2.sheet_view.rightToLeft)
print('fusions            :', len(ws2.merged_cells.ranges), sorted(str(x) for x in ws2.merged_cells.ranges))
print('en-tete de la grille (col B..I) :')
for col in range(2, 10):
    print('   ', get_column_letter(col), '=', ws2.cell(row=LIGNE_EN_TETE, column=col).value)
print('grille (jour -> cases) :')
for i, jour in enumerate(JOURS):
    ligne = LIGNE_EN_TETE + 1 + i
    cases = []
    for col in range(2, 10):
        v = ws2.cell(row=ligne, column=col).value
        if v is None:
            continue
        txt = str(v).replace('\n', ' / ')
        cases.append(get_column_letter(col) + ':' + txt)
    print('   %-9s %s' % (jour, ' | '.join(cases) if cases else '(aucune seance)'))
