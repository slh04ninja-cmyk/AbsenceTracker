# -*- coding: utf-8 -*-
"""Genere les emplois du temps TCSF-1, TCSF-2, TCSF-3 en xlsx (mise en forme arabe, feuille RTL)
a partir de _edt_classes.json. Meme presentation que tab_service_TCSLHF1.xlsx."""
import io, json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

EDT = json.load(io.open('_edt_classes.json', encoding='utf-8'))
JOURS = {'1': 'الإثنين', '2': 'الثلاثاء', '3': 'الأربعاء', '4': 'الخميس', '5': 'الجمعة', '6': 'السبت'}
CRENEAUX = ['08:00-09:00', '09:00-10:00', '10:00-11:00', '11:00-12:00',
            '14:00-15:00', '15:00-16:00', '16:00-17:00', '17:00-18:00']
MATIERE_AR = {'Maths': 'الرياضيات', 'PC': 'الفيزياء والكيمياء', 'SVT': 'علوم الحياة والأرض',
              'Français': 'اللغة الفرنسية', 'Anglais': 'اللغة الإنجليزية', 'Arabe': 'اللغة العربية',
              'EPS': 'التربية البدنية', 'Informatique': 'المعلوميات', 'Philo': 'الفلسفة',
              'Educ. islamique': 'التربية الإسلامية', 'Hist-Géo': 'التاريخ والجغرافيا'}
PROFS = {'Maths': ('الرياضيات', 'أيوب الكمرة'), 'PC': ('الفيزياء والكيمياء', 'غزالي صالح'),
         'SVT': ('علوم الحياة والأرض', 'كمال الوردي'), 'Français': ('اللغة الفرنسية', 'سامية الحاضي'),
         'Anglais': ('اللغة الإنجليزية', 'هشام أجامي'), 'Arabe': ('اللغة العربية', 'المهدي الصلحي'),
         'EPS': ('التربية البدنية', 'يسرى البوسعيدي'), 'Informatique': ('المعلوميات', 'سكينة الرازي'),
         'Philo': ('الفلسفة', 'زكية المندريلي'), 'Educ. islamique': ('التربية الإسلامية', 'محمد خليفي'),
         'Hist-Géo': ('التاريخ والجغرافيا', 'ياسين القامة')}
ORDRE_MATIERES = ['Maths', 'PC', 'SVT', 'Français', 'Anglais', 'Arabe', 'EPS',
                  'Informatique', 'Philo', 'Educ. islamique', 'Hist-Géo']

FIN = Side(style='thin', color='000000')
BORD = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)
GRIS = PatternFill('solid', fgColor='F2F2F2')
GRIS_FONCE = PatternFill('solid', fgColor='E7E6E6')
CENTRE = Alignment(horizontal='center', vertical='center', wrap_text=True)
LARG = 9

for classe, parJour in EDT.items():
    wb = Workbook()
    ws = wb.active
    ws.title = classe
    ws.sheet_view.rightToLeft = True

    def entete(idx, texte, gras=True, taille=11, fond=None):
        ws.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=LARG)
        c = ws.cell(row=idx, column=1, value=texte)
        c.font = Font(bold=gras, size=taille)
        c.alignment = CENTRE
        if fond is not None:
            for col in range(1, LARG + 1):
                ws.cell(row=idx, column=col).fill = fond
        ws.row_dimensions[idx].height = 20

    r = 1
    entete(r, 'المملكة المغربية', True, 12); r += 1
    entete(r, 'وزارة التربية الوطنية والتعليم الأولي والرياضة', False, 10); r += 1
    entete(r, 'المديرية الإقليمية – سطات', False, 10); r += 1
    entete(r, 'الثانوية القصبية الإعدادية – بوعوان', True, 11); r += 1
    entete(r, 'الموسم الدراسي 2026-2027', False, 10); r += 1
    entete(r, 'جدول حصص القسم : ' + classe, True, 13, GRIS); r += 1
    entete(r, 'EDT for FET timetables', False, 8); r += 2

    ligne_entete = r
    ws.cell(row=r, column=1, value='اليوم / الحصة').font = Font(bold=True, size=10)
    ws.cell(row=r, column=1).alignment = CENTRE
    ws.cell(row=r, column=1).fill = GRIS_FONCE
    for i, cr in enumerate(CRENEAUX):
        c = ws.cell(row=r, column=2 + i, value=cr)
        c.font = Font(bold=True, size=9)
        c.alignment = CENTRE
        c.fill = GRIS_FONCE
    for col in range(1, LARG + 1):
        ws.cell(row=r, column=col).border = BORD
    ws.row_dimensions[r].height = 24
    r += 1

    for j in ['1', '2', '3', '4', '5', '6']:
        c = ws.cell(row=r, column=1, value=JOURS[j])
        c.font = Font(bold=True, size=11)
        c.alignment = CENTRE
        c.fill = GRIS
        c.border = BORD
        for col in range(2, LARG + 1):
            cc = ws.cell(row=r, column=col)
            cc.border = BORD
            cc.alignment = CENTRE
            cc.font = Font(size=10)
        for creneau in parJour.get(j, []):
            col = 2 + creneau['debut']
            cellule = ws.cell(row=r, column=col)
            cellule.value = MATIERE_AR.get(creneau['matiere'], creneau['matiere'])
            if creneau['duree'] > 1:
                ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + creneau['duree'] - 1)
        ws.row_dimensions[r].height = 38
        r += 1

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=LARG)
    c = ws.cell(row=r, column=1, value='لائحة الأساتذة المسندين حسب المواد')
    c.font = Font(bold=True, size=11)
    c.alignment = CENTRE
    for col in range(1, LARG + 1):
        ws.cell(row=r, column=col).fill = GRIS
        ws.cell(row=r, column=col).border = BORD
    r += 1
    paires = [1, 4, 7]
    for ligne_idx in range(4):
        for i, col_dep in enumerate(paires):
            k = ligne_idx + 4 * i
            if k >= len(ORDRE_MATIERES):
                continue
            mat = ORDRE_MATIERES[k]
            m_ar, prof = PROFS[mat]
            cm = ws.cell(row=r, column=col_dep, value=m_ar)
            cp = ws.cell(row=r, column=col_dep + 1, value=prof)
            for cc in (cm, cp):
                cc.border = BORD
                cc.alignment = CENTRE
                cc.font = Font(size=10, bold=(cc is cm))
        ws.row_dimensions[r].height = 20
        r += 1

    ws.column_dimensions['A'].width = 16
    for i in range(2, LARG + 1):
        ws.column_dimensions[get_column_letter(i)].width = 13
    ws.freeze_panes = ws.cell(row=ligne_entete + 1, column=2)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    nom = 'emploi_du_temps_%s.xlsx' % classe
    wb.save(nom)
    total = sum(sum(x['duree'] for x in parJour.get(j, [])) for j in '12345')
    print('%-26s %s h | samedi : %s' % (nom, total, parJour.get('6', [])))
print('termine')
